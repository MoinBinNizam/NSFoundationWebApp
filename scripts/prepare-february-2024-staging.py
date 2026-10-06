"""Create a reviewable February 2024 member-by-member import manifest.

The source workbook is never changed.  Its explicit paid-in-January labels are
linked to the January advance ledger rather than duplicated as February cash.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

SOURCE = Path("docs/import-historical-evidence/Month_February_2024.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/february-2024-staging")
MONTH = "2024-02"
RATE = 500.0


def value(item: Any) -> str:
    if item is None:
        return ""
    if isinstance(item, (date, datetime)):
        return item.strftime("%Y-%m-%d")
    return str(item).strip()


def number(item: Any) -> float:
    try:
        return float(item)
    except (TypeError, ValueError):
        return 0.0


def gateway(item: Any) -> str:
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "cash": "CASH", "self": "CASH"}.get(value(item).lower(), "")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["February, 2024"]
    ledgers, payments, allocations = [], [], []
    for row in range(3, 38):
        member_id = value(ws.cell(row, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name, shares = value(ws.cell(row, 2).value), number(ws.cell(row, 3).value)
        raw_amount, source_due, source_penalty = ws.cell(row, 4).value, number(ws.cell(row, 8).value), number(ws.cell(row, 9).value)
        comment = value(ws.cell(row, 10).value)
        principal_due = shares * RATE
        paid_in_january = value(raw_amount).lower() == "paid" and ("january" in comment.lower() or " jan" in comment.lower())
        received = number(raw_amount)
        # NSF032's source comment explicitly says the February receipt paid its
        # own January arrear plus February. It is never used for another member.
        january_arrear_payment = received - principal_due if member_id == "NSF032" else 0.0
        feb_principal_paid = principal_due if paid_in_january else min(principal_due, max(0.0, received - january_arrear_payment))
        penalty_due = source_penalty if member_id == "NSF013" else 0.0
        ref = f"HIST-2024-02-R{row:02d}"
        ledgers.append({"source_ref": ref, "member_id": member_id, "member_name": name, "month": MONTH, "share_count": shares, "principal_due": principal_due, "principal_paid": feb_principal_paid, "penalty_due": penalty_due, "penalty_paid": 0.0, "advance_applied": feb_principal_paid if paid_in_january else 0.0, "excess_advance": 0.0, "status": "PAID" if feb_principal_paid >= principal_due and penalty_due == 0 else "DUE", "source_row": row, "comment": comment})
        if received <= 0:
            continue
        payment_ref = f"{ref}-P01"
        payments.append({"payment_ref": payment_ref, "source_ref": ref, "member_id": member_id, "member_name": name, "payment_date": value(ws.cell(row, 7).value), "total_amount": received, "gateway": gateway(ws.cell(row, 5).value), "receiver_source": value(ws.cell(row, 6).value), "comment": comment})
        if january_arrear_payment:
            allocations.append({"payment_ref": payment_ref, "member_id": member_id, "target_month": "2024-01", "allocation_type": "PRINCIPAL", "amount": january_arrear_payment, "source_ref": ref})
        if feb_principal_paid:
            allocations.append({"payment_ref": payment_ref, "member_id": member_id, "target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": feb_principal_paid, "source_ref": ref})
    manifest = {"ledgers": ledgers, "payments": payments, "allocations": allocations}
    (OUT / "staging-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    summary = {"members": len(ledgers), "payments": len(payments), "allocations": len(allocations), "cash_received": sum(row["total_amount"] for row in payments), "feb_principal_due": sum(row["principal_due"] - row["principal_paid"] for row in ledgers), "feb_penalty_due": sum(row["penalty_due"] for row in ledgers), "january_arrear_collected": sum(row["amount"] for row in allocations if row["target_month"] == "2024-01")}
    (OUT / "verification-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
