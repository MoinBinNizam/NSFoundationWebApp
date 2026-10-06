"""Create a reviewable August 2024 member-by-member import manifest.

The source workbook is never changed.
Advance paid in May for NSF031 is applied to August.
August penalties are waived due to the July Revolution.
Cross-month FIFO arrear allocations for NSF014, NSF025, NSF033, and NSF034 are resolved cleanly.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

SOURCE = Path("docs/import-historical-evidence/Month_August_2024.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/august-2024-staging")
MONTH = "2024-08"
RATE = 500.0


def value(item: Any) -> str:
    if item is None:
        return ""
    if isinstance(item, (date, datetime)):
        return item.strftime("%Y-%m-%d")
    return str(item).strip()


def number(item: Any) -> float:
    try:
        return float(item)
    except (TypeError, ValueError):
        return 0.0


def gateway(item: Any) -> str:
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "cash": "CASH", "self": "CASH"}.get(value(item).lower(), "")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["August, 2024"]
    ledgers, payments, allocations = [], [], []

    for row in range(3, 38):
        member_id = value(ws.cell(row, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name, shares = value(ws.cell(row, 2).value), number(ws.cell(row, 3).value)
        raw_amount = ws.cell(row, 4).value
        source_due, source_penalty = number(ws.cell(row, 8).value), number(ws.cell(row, 9).value)
        comment = value(ws.cell(row, 10).value)
        principal_due = shares * RATE
        paid_in_advance = value(raw_amount).lower() == "paid" and ("advanced" in comment.lower() or "may" in comment.lower())
        received = number(raw_amount)

        custom_allocations = []
        august_principal_paid = 0.0
        penalty_due = 0.0
        penalty_paid = 0.0
        excess_advance = 0.0

        if member_id == "NSF013":
            # Paid ৳2,000 on 2024-08-30: ৳1,000 August principal + ৳1,000 excess advance
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "ADVANCE", "amount": 1000.0})
            august_principal_paid = 1000.0
            excess_advance = 1000.0
        elif member_id == "NSF014":
            # Paid ৳6,120 on 2024-08-12:
            # ৳3,000 July principal + ৳120 July penalty + ৳3,000 August principal
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PRINCIPAL", "amount": 3000.0})
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PENALTY", "amount": 120.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 3000.0})
            august_principal_paid = 3000.0
        elif member_id == "NSF025":
            # Paid ৳2,010 on 2024-08-20:
            # ৳1,000 July principal + ৳10 July penalty + ৳1,000 August principal
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PENALTY", "amount": 10.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            august_principal_paid = 1000.0
        elif member_id == "NSF033":
            # Paid ৳5,160 on 2024-08-28:
            # ৳1,000 April principal + ৳40 April penalty
            # ৳1,000 May principal + ৳40 May penalty
            # ৳1,000 June principal + ৳40 June penalty
            # ৳1,000 July principal + ৳40 July penalty
            # ৳1,000 August principal
            custom_allocations.append({"target_month": "2024-04", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-04", "allocation_type": "PENALTY", "amount": 40.0})
            custom_allocations.append({"target_month": "2024-05", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-05", "allocation_type": "PENALTY", "amount": 40.0})
            custom_allocations.append({"target_month": "2024-06", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-06", "allocation_type": "PENALTY", "amount": 40.0})
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-07", "allocation_type": "PENALTY", "amount": 40.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            august_principal_paid = 1000.0
        elif member_id == "NSF034":
            # Paid ৳3,000 on 2024-08-30. Under FIFO, settles March 2024 principal arrear
            custom_allocations.append({"target_month": "2024-03", "allocation_type": "PRINCIPAL", "amount": 3000.0})
            august_principal_paid = 0.0
        elif paid_in_advance:
            # Paid via May advance (NSF031: ৳500)
            august_principal_paid = principal_due
        elif received > 0:
            august_principal_paid = min(principal_due, received)
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": august_principal_paid})
        else:
            august_principal_paid = 0.0

        ref = f"HIST-2024-08-R{row:02d}"
        status = "PAID" if (august_principal_paid >= principal_due and penalty_due == 0) else ("PARTIAL" if august_principal_paid > 0 else "DUE")

        ledgers.append({
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "month": MONTH,
            "share_count": shares,
            "principal_due": principal_due,
            "principal_paid": august_principal_paid,
            "penalty_due": penalty_due,
            "penalty_paid": penalty_paid,
            "advance_applied": august_principal_paid if paid_in_advance else 0.0,
            "excess_advance": excess_advance,
            "status": status,
            "source_row": row,
            "comment": comment
        })

        if received <= 0:
            continue

        payment_ref = f"{ref}-P01"
        payments.append({
            "payment_ref": payment_ref,
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "payment_date": value(ws.cell(row, 7).value),
            "total_amount": received,
            "gateway": gateway(ws.cell(row, 5).value),
            "receiver_source": value(ws.cell(row, 6).value),
            "comment": comment
        })

        for alloc in custom_allocations:
            allocations.append({
                "payment_ref": payment_ref,
                "member_id": member_id,
                "target_month": alloc["target_month"],
                "allocation_type": alloc["allocation_type"],
                "amount": alloc["amount"],
                "source_ref": ref
            })

    manifest = {"ledgers": ledgers, "payments": payments, "allocations": allocations}
    (OUT / "staging-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    summary = {
        "members": len(ledgers),
        "payments": len(payments),
        "allocations": len(allocations),
        "cash_received": sum(row["total_amount"] for row in payments),
        "august_principal_due": sum(row["principal_due"] - row["principal_paid"] for row in ledgers),
        "prior_principal_arrears_collected": sum(row["amount"] for row in allocations if row["target_month"] < MONTH and row["allocation_type"] == "PRINCIPAL"),
        "prior_penalties_collected": sum(row["amount"] for row in allocations if row["target_month"] < MONTH and row["allocation_type"] == "PENALTY"),
        "advances_applied": sum(row["advance_applied"] for row in ledgers),
    }
    (OUT / "verification-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
