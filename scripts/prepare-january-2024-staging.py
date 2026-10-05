"""Prepare a review-only January 2024 staging package from the authoritative sheet.

This script never connects to MongoDB. It creates member-specific payment,
monthly-ledger, and allocation manifests ready for approval before a future
staging-database posting operation.
"""

from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


SOURCE = Path("docs/import-historical-evidence/Month_January_2024.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/january-2024-staging")
MONTH = "2024-01"
SHARE_RATE = 500.0

# These are explicitly stated in the authoritative source comments. Values are
# allocations for that member only; no amount can be offset against another ID.
ADVANCES = {
    "NSF023": [("2024-02", 500.0)],
    "NSF030": [("2024-02", 1000.0)],
    "NSF031": [("2024-02", 500.0), ("2024-03", 500.0), ("2024-04", 500.0)],
}
SPLIT_RECEIPTS = {
    "NSF001": [("2024-01-02", 1000.0), ("2024-01-21", 1500.0)],
    "NSF002": [("2024-01-02", 1000.0), ("2024-01-23", 1000.0)],
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def amount(value: Any) -> float:
    return float(value or 0)


def gateway(value: Any) -> str:
    raw = text(value).lower()
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "self": "CASH", "cash": "CASH"}.get(raw, "")


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else [])
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["January, 2024"]
    ledger_rows: list[dict[str, Any]] = []
    payment_rows: list[dict[str, Any]] = []
    allocation_rows: list[dict[str, Any]] = []

    for row_number in range(3, ws.max_row + 1):
        member_id = text(ws.cell(row_number, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name = text(ws.cell(row_number, 2).value)
        shares = amount(ws.cell(row_number, 3).value)
        received = amount(ws.cell(row_number, 4).value)
        source_date = text(ws.cell(row_number, 7).value)
        comment = text(ws.cell(row_number, 10).value)
        principal_due = shares * SHARE_RATE
        source_ref = f"HIST-2024-01-R{row_number:02d}"
        advance_total = sum(value for _, value in ADVANCES.get(member_id, []))
        principal_paid_january = min(received, principal_due)
        due = max(0.0, principal_due - principal_paid_january)
        ledger_rows.append({
            "source_ref": source_ref, "member_id": member_id, "member_name": name, "month": MONTH,
            "share_count": f"{shares:.0f}", "principal_due": f"{principal_due:.2f}",
            "principal_paid": f"{principal_paid_january:.2f}", "penalty_due": "0.00", "penalty_paid": "0.00",
            "advance_applied": "0.00", "excess_advance": f"{advance_total:.2f}", "status": "PAID" if due == 0 else "DUE",
            "source_row": row_number, "comment": comment,
        })

        if received <= 0:
            continue
        receipt_parts = SPLIT_RECEIPTS.get(member_id, [(source_date, received)])
        if round(sum(part_amount for _, part_amount in receipt_parts), 2) != round(received, 2):
            raise ValueError(f"Receipt split does not equal source amount for {member_id}")
        remaining_january = principal_due
        for index, (payment_date, part_amount) in enumerate(receipt_parts, start=1):
            payment_ref = f"{source_ref}-P{index:02d}"
            principal_portion = min(remaining_january, part_amount)
            remaining_january -= principal_portion
            payment_rows.append({
                "payment_ref": payment_ref, "source_ref": source_ref, "member_id": member_id, "member_name": name,
                "payment_date": payment_date, "total_amount": f"{part_amount:.2f}", "gateway": gateway(ws.cell(row_number, 5).value),
                "receiver_source": text(ws.cell(row_number, 6).value), "penalty_amount": "0.00", "cashout_charge": "0.00",
                "transaction_reference": "", "comment": comment, "posting_decision": "READY_FOR_STAGING_AFTER_MEMBER_AND_CUSTODY_MAPPING",
            })
            if principal_portion:
                allocation_rows.append({"payment_ref": payment_ref, "member_id": member_id, "target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": f"{principal_portion:.2f}", "source_ref": source_ref})

        # Advance allocation is added after the current month allocation. It is
        # attached to the original source payment and never to another member.
        if ADVANCES.get(member_id):
            payment_ref = f"{source_ref}-P{len(receipt_parts):02d}"
            for target_month, advance_amount in ADVANCES[member_id]:
                allocation_rows.append({"payment_ref": payment_ref, "member_id": member_id, "target_month": target_month, "allocation_type": "ADVANCE", "amount": f"{advance_amount:.2f}", "source_ref": source_ref})
                ledger_rows.append({
                    "source_ref": f"{source_ref}-ADV-{target_month}", "member_id": member_id, "member_name": name, "month": target_month,
                    "share_count": f"{shares:.0f}", "principal_due": f"{advance_amount:.2f}", "principal_paid": f"{advance_amount:.2f}",
                    "penalty_due": "0.00", "penalty_paid": "0.00", "advance_applied": f"{advance_amount:.2f}", "excess_advance": "0.00",
                    "status": "PAID", "source_row": row_number, "comment": f"Advance from {MONTH}: {comment}",
                })

    write_csv(OUT / "monthly-ledger-manifest.csv", ledger_rows)
    write_csv(OUT / "payment-manifest.csv", payment_rows)
    write_csv(OUT / "payment-allocation-manifest.csv", allocation_rows)
    (OUT / "staging-manifest.json").write_text(json.dumps({"ledgers": ledger_rows, "payments": payment_rows, "allocations": allocation_rows}, indent=2), encoding="utf-8")
    payment_total = sum(float(row["total_amount"]) for row in payment_rows)
    due_total = sum(float(row["principal_due"]) - float(row["principal_paid"]) for row in ledger_rows)
    advance_total = sum(float(row["amount"]) for row in allocation_rows if row["allocation_type"] == "ADVANCE")
    report = f"""# January 2024 staging package\n\nThis package was prepared from the authoritative January 2024 workbook. It is review-only and has not been posted to MongoDB.\n\n| Control | Value |\n|---|---:|\n| Member ledgers | {len(ledger_rows)} |\n| Historical payment records | {len(payment_rows)} |\n| Payment allocation records | {len(allocation_rows)} |\n| Payment total | ৳{payment_total:,.2f} |\n| January principal due remaining | ৳{due_total:,.2f} |\n| Member-specific future advance | ৳{advance_total:,.2f} |\n\n## Rules applied\n\n- NSF032, NSF034, and NSF035 retain their own January dues totalling ৳4,500.\n- Advances remain under the paying member only: NSF023 ৳500; NSF030 ৳1,000; NSF031 ৳1,500.\n- NSF001 and NSF002 are represented as two historical receipts each because the source comments provide separate payment dates.\n- January has no source penalty or cash-out collection.\n- `Self` uses Cash custody; Bank, bKash, and Nagad retain their source gateway.\n\n## Posting gate\n\nBefore a staging-database posting run, map every manifest member ID to the registered Member record and each source receiver/gateway pair to an active historical custody account. Preserve the `source_ref` in notes/audit evidence and do not generate or invent external transaction IDs.\n"""
    (OUT / "README.md").write_text(report, encoding="utf-8")
    print(f"Prepared {len(ledger_rows)} ledgers, {len(payment_rows)} payments, and {len(allocation_rows)} allocations in {OUT}")


if __name__ == "__main__":
    main()
