"""Create a reviewable October 2024 member-by-member import manifest.

The source workbook is never changed.
Cross-month FIFO arrear and advance allocations for NSF005, NSF006, NSF008, NSF012, NSF018, NSF033, NSF034, and NSF035 are resolved cleanly.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

SOURCE = Path("docs/import-historical-evidence/Month_October_2024.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/october-2024-staging")
MONTH = "2024-10"
RATE = 500.0


def value(item: Any) -> str:
    if item is None:
        return ""
    if isinstance(item, (date, datetime)):
        return item.strftime("%Y-%m-%d")
    return str(item).strip()


def number(item: Any) -> float:
    try:
        return float(item)
    except (TypeError, ValueError):
        return 0.0


def gateway(item: Any) -> str:
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "cash": "CASH", "self": "CASH"}.get(value(item).lower(), "")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["Oct, 2024"]
    ledgers, payments, allocations = [], [], []

    for row in range(3, 38):
        member_id = value(ws.cell(row, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name, shares = value(ws.cell(row, 2).value), number(ws.cell(row, 3).value)
        raw_amount = ws.cell(row, 4).value
        source_due, source_penalty = number(ws.cell(row, 8).value), number(ws.cell(row, 9).value)
        comment = value(ws.cell(row, 10).value)
        principal_due = shares * RATE
        received = number(raw_amount)

        custom_allocations = []
        oct_principal_paid = 0.0
        penalty_due = 0.0
        penalty_paid = 0.0
        advance_applied = 0.0
        excess_advance = 0.0

        if member_id == "NSF005":
            # Paid in advance in September 2024 (৳1,000 October principal + ৳30 excess advance)
            oct_principal_paid = 1000.0
            advance_applied = 1000.0
            excess_advance = 30.0
        elif member_id == "NSF006":
            # Paid ৳1,020 on 2024-10-03: ৳500 Sept principal + ৳20 Sept penalty + ৳500 Oct principal
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PENALTY", "amount": 20.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            oct_principal_paid = 500.0
        elif member_id == "NSF008":
            # Paid ৳530 on 2024-10-16: ৳10 Sept penalty (half of 20 due) + ৳500 Oct principal + ৳20 Oct penalty
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PENALTY", "amount": 10.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PENALTY", "amount": 20.0})
            oct_principal_paid = 500.0
            penalty_due = 20.0
            penalty_paid = 20.0
        elif member_id == "NSF012":
            # Paid ৳3,100 on 2024-10-15:
            # ৳30 June remaining penalty + ৳1,000 August principal + ৳1,000 Sept principal + ৳70 Sept penalty + ৳1,000 Oct principal
            custom_allocations.append({"target_month": "2024-06", "allocation_type": "PENALTY", "amount": 30.0})
            custom_allocations.append({"target_month": "2024-08", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PENALTY", "amount": 70.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            oct_principal_paid = 1000.0
        elif member_id == "NSF013":
            # Paid ৳1,000 on 2024-10-17: ৳1,000 Oct principal, ৳40 Oct penalty unpaid
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            oct_principal_paid = 1000.0
            penalty_due = 40.0
            penalty_paid = 0.0
        elif member_id == "NSF018":
            # Paid ৳1,040 on 2024-10-07: ৳40 April penalty + ৳1,000 Oct principal
            custom_allocations.append({"target_month": "2024-04", "allocation_type": "PENALTY", "amount": 40.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            oct_principal_paid = 1000.0
        elif member_id == "NSF026":
            # Paid ৳520 on 2024-10-24 (late): ৳500 Oct principal + ৳20 Oct penalty
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PENALTY", "amount": 20.0})
            oct_principal_paid = 500.0
            penalty_due = 20.0
            penalty_paid = 20.0
        elif member_id == "NSF028":
            # Paid ৳1,010 on 2024-10-09: ৳1,000 Oct principal + ৳10 advance (cashout charge)
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "ADVANCE", "amount": 10.0})
            oct_principal_paid = 1000.0
            excess_advance = 10.0
        elif member_id == "NSF033":
            # Paid ৳2,020 on 2024-10-17: ৳1,000 Sept principal + ৳1,000 Oct principal + ৳20 Sept penalty
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PENALTY", "amount": 20.0})
            oct_principal_paid = 1000.0
            penalty_due = 40.0
            penalty_paid = 0.0
        elif member_id == "NSF034":
            # Paid ৳6,240 on 2024-10-27: ৳3,000 Sept principal + ৳120 Sept penalty + ৳3,000 Oct principal + ৳120 Oct penalty
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PRINCIPAL", "amount": 3000.0})
            custom_allocations.append({"target_month": "2024-09", "allocation_type": "PENALTY", "amount": 120.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 3000.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PENALTY", "amount": 120.0})
            oct_principal_paid = 3000.0
            penalty_due = 120.0
            penalty_paid = 120.0
        elif member_id == "NSF035":
            # Paid ৳6,000 on 2024-10-07: Joined today, paid all 12 months (Jan-Dec 2024 @ ৳500/mo) at once
            for m in range(1, 11):
                custom_allocations.append({"target_month": f"2024-{m:02d}", "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": "2024-11", "allocation_type": "ADVANCE", "amount": 500.0})
            custom_allocations.append({"target_month": "2024-12", "allocation_type": "ADVANCE", "amount": 500.0})
            oct_principal_paid = 500.0
        elif received > 0:
            oct_principal_paid = min(principal_due, received)
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": oct_principal_paid})
        else:
            oct_principal_paid = 0.0
            penalty_due = source_penalty if member_id in ("NSF007", "NSF015", "NSF016", "NSF019", "NSF020", "NSF022", "NSF023", "NSF025", "NSF027", "NSF031", "NSF032") else (shares * 20.0)

        ref = f"HIST-2024-10-R{row:02d}"

        ledgers.append({
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "month": MONTH,
            "share_count": int(shares),
            "principal_due": principal_due,
            "principal_paid": oct_principal_paid,
            "penalty_due": penalty_due,
            "penalty_paid": penalty_paid,
            "advance_applied": advance_applied,
            "excess_advance": excess_advance,
            "status": "PAID" if oct_principal_paid >= principal_due and penalty_due == penalty_paid else ("PARTIAL" if (oct_principal_paid > 0 or penalty_paid > 0) else "DUE"),
            "comment": comment,
        })

        if received > 0:
            pref = f"PAY-2024-10-R{row:02d}"
            payments.append({
                "payment_ref": pref,
                "source_ref": ref,
                "member_id": member_id,
                "member_name": name,
                "payment_date": value(ws.cell(row, 7).value),
                "total_amount": received,
                "gateway": gateway(ws.cell(row, 5).value),
                "receiver_source": value(ws.cell(row, 6).value).capitalize(),
                "comment": comment,
            })
            for alloc in custom_allocations:
                allocations.append({
                    "payment_ref": pref,
                    "member_id": member_id,
                    "target_month": alloc["target_month"],
                    "allocation_type": alloc["allocation_type"],
                    "amount": alloc["amount"],
                    "source_ref": ref,
                })

    manifest = {"ledgers": ledgers, "payments": payments, "allocations": allocations}
    (OUT / "staging-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    verification = {
        "month": MONTH,
        "members": len(ledgers),
        "totalShares": sum(l["share_count"] for l in ledgers),
        "paymentsCount": len(payments),
        "paymentsTotal": sum(p["total_amount"] for p in payments),
        "allocationsCount": len(allocations),
        "allocationsTotal": sum(a["amount"] for a in allocations),
        "byReceiver": {
            "Moin": sum(p["total_amount"] for p in payments if "moin" in p["receiver_source"].lower()),
            "Samrat": sum(p["total_amount"] for p in payments if "samrat" in p["receiver_source"].lower()),
        },
    }
    (OUT / "verification-summary.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
    print(json.dumps(verification, indent=2))


if __name__ == "__main__":
    main()
