"""Create a reviewable January 2025 member-by-member import manifest.

The source workbook is never changed.
Cross-month FIFO arrear settlements for NSF012 and NSF027, and advances for NSF029 and NSF035 are resolved cleanly.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

SOURCE = Path("docs/import-historical-evidence/Month_January_2025.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/january-2025-staging")
MONTH = "2025-01"
RATE = 500.0


def value(item: Any) -> str:
    if item is None:
        return ""
    if isinstance(item, (date, datetime)):
        return item.strftime("%Y-%m-%d")
    return str(item).strip()


def number(item: Any) -> float:
    try:
        return float(item)
    except (TypeError, ValueError):
        return 0.0


def gateway(item: Any) -> str:
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "cash": "CASH", "self": "CASH"}.get(value(item).lower(), "")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["January 25"]
    ledgers, payments, allocations = [], [], []

    for row in range(3, 38):
        member_id = value(ws.cell(row, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name, shares = value(ws.cell(row, 2).value), number(ws.cell(row, 3).value)
        raw_paid = ws.cell(row, 5).value
        source_due_pen = number(ws.cell(row, 7).value)
        comment = ""
        principal_due = shares * RATE
        received = number(raw_paid)

        custom_allocations = []
        jan_principal_paid = 0.0
        penalty_due = 0.0
        penalty_paid = 0.0
        advance_applied = 0.0
        excess_advance = 0.0

        if member_id == "NSF012":
            # Paid ৳3,200 on 2025-01-07:
            # ৳1,000 Nov principal + ৳1,000 Dec principal + ৳1,000 Jan principal + ৳80 penalty + ৳120 advance (45 cashout + 75 balance)
            custom_allocations.append({"target_month": "2024-11", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-12", "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 1000.0})
            custom_allocations.append({"target_month": "2024-11", "allocation_type": "PENALTY", "amount": 80.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "ADVANCE", "amount": 120.0})
            jan_principal_paid = 1000.0
            excess_advance = 120.0
        elif member_id == "NSF027":
            # Paid ৳1,120 on 2025-01-06: ৳500 Dec past due + ৳120 Dec penalty + ৳500 Jan principal
            custom_allocations.append({"target_month": "2024-12", "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": "2024-12", "allocation_type": "PENALTY", "amount": 120.0})
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            jan_principal_paid = 500.0
        elif member_id == "NSF029":
            # Paid ৳1,000 on 2025-01-06: ৳500 Jan principal + ৳500 Feb advance
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            custom_allocations.append({"target_month": "2025-02", "allocation_type": "ADVANCE", "amount": 500.0})
            jan_principal_paid = 500.0
        elif member_id == "NSF035":
            # Paid ৳6,000 on 2025-01-27: ৳500 Jan principal + ৳500 for each month Feb-Dec 2025 advance (11 * 500 = 5500)
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": 500.0})
            for m in range(2, 13):
                custom_allocations.append({"target_month": f"2025-{m:02d}", "allocation_type": "ADVANCE", "amount": 500.0})
            jan_principal_paid = 500.0
        elif received > 0:
            jan_principal_paid = min(principal_due, received)
            custom_allocations.append({"target_month": MONTH, "allocation_type": "PRINCIPAL", "amount": jan_principal_paid})
        else:
            jan_principal_paid = 0.0
            penalty_due = source_due_pen or (shares * 20.0)

        ref = f"HIST-2025-01-R{row:02d}"

        ledgers.append({
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "month": MONTH,
            "share_count": int(shares),
            "principal_due": principal_due,
            "principal_paid": jan_principal_paid,
            "penalty_due": penalty_due,
            "penalty_paid": penalty_paid,
            "advance_applied": advance_applied,
            "excess_advance": excess_advance,
            "status": "PAID" if jan_principal_paid >= principal_due and penalty_due == penalty_paid else ("PARTIAL" if (jan_principal_paid > 0 or penalty_paid > 0) else "DUE"),
            "comment": comment,
        })

        if received > 0:
            pref = f"PAY-2025-01-R{row:02d}"
            payments.append({
                "payment_ref": pref,
                "source_ref": ref,
                "member_id": member_id,
                "member_name": name,
                "payment_date": value(ws.cell(row, 6).value),
                "total_amount": received,
                "gateway": gateway(ws.cell(row, 12).value),
                "receiver_source": value(ws.cell(row, 13).value).capitalize(),
                "comment": comment,
            })
            for alloc in custom_allocations:
                allocations.append({
                    "payment_ref": pref,
                    "member_id": member_id,
                    "target_month": alloc["target_month"],
                    "allocation_type": alloc["allocation_type"],
                    "amount": alloc["amount"],
                    "source_ref": ref,
                })

    manifest = {"ledgers": ledgers, "payments": payments, "allocations": allocations}
    (OUT / "staging-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    verification = {
        "month": MONTH,
        "members": len(ledgers),
        "totalShares": sum(l["share_count"] for l in ledgers),
        "paymentsCount": len(payments),
        "paymentsTotal": sum(p["total_amount"] for p in payments),
        "allocationsCount": len(allocations),
        "allocationsTotal": sum(a["amount"] for a in allocations),
        "byReceiver": {
            "Moin": sum(p["total_amount"] for p in payments if "moin" in p["receiver_source"].lower()),
            "Samrat": sum(p["total_amount"] for p in payments if "samrat" in p["receiver_source"].lower()),
        },
        "byGateway": {
            "BANK": sum(p["total_amount"] for p in payments if p["gateway"] == "BANK"),
            "BKASH": sum(p["total_amount"] for p in payments if p["gateway"] == "BKASH"),
            "NAGAD": sum(p["total_amount"] for p in payments if p["gateway"] == "NAGAD"),
        },
    }
    (OUT / "verification-summary.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
    print(json.dumps(verification, indent=2))


if __name__ == "__main__":
    main()
