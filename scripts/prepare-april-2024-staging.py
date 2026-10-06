"""Create a reviewable April 2024 member-by-member import manifest.

The source workbook is never changed. Explicit paid-in-January labels are
linked to the January advance ledger rather than duplicated as April cash.
Prior March arrears and penalties for NSF024 and NSF027 are allocated accurately.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

SOURCE = Path("docs/import-historical-evidence/Month_April_2024.xlsx")
OUT = Path("docs/import-historical-evidence/normalized-manifests/april-2024-staging")
MONTH = "2024-04"
RATE = 500.0


def value(item: Any) -> str:
    if item is None:
        return ""
    if isinstance(item, (date, datetime)):
        # Normalize date typo in workbook (e.g., 2023-04-05 -> 2024-04-05)
        d_str = item.strftime("%Y-%m-%d")
        if d_str.startswith("2023-04"):
            return d_str.replace("2023-04", "2024-04")
        return d_str
    return str(item).strip()


def number(item: Any) -> float:
    try:
        return float(item)
    except (TypeError, ValueError):
        return 0.0


def gateway(item: Any) -> str:
    return {"bkash": "BKASH", "nagad": "NAGAD", "bank": "BANK", "cash": "CASH", "self": "CASH"}.get(value(item).lower(), "")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SOURCE, data_only=True)["April, 2024"]
    ledgers, payments, allocations = [], [], []

    for row in range(3, 38):
        member_id = value(ws.cell(row, 1).value).upper()
        if not member_id.startswith("NSF"):
            continue
        name, shares = value(ws.cell(row, 2).value), number(ws.cell(row, 3).value)
        raw_amount = ws.cell(row, 4).value
        source_due, source_penalty = number(ws.cell(row, 8).value), number(ws.cell(row, 9).value)
        comment = value(ws.cell(row, 10).value)
        principal_due = shares * RATE
        paid_in_january = value(raw_amount).lower() == "paid" and ("january" in comment.lower() or " jan" in comment.lower())
        received = number(raw_amount)

        # Cross-month arrear and penalty splits
        march_principal_arrear = 0.0
        march_penalty_arrear = 0.0
        april_penalty_paid = 0.0

        if member_id == "NSF024":
            # Paid ৳540: ৳20 March penalty + ৳500 April principal + ৳20 April penalty
            march_penalty_arrear = 20.0
            april_principal_paid = 500.0
            april_penalty_due = 20.0
            april_penalty_paid = 20.0
        elif member_id == "NSF027":
            # Paid ৳2,020: ৳1,000 March principal + ৳20 March penalty + ৳1,000 April principal
            march_principal_arrear = 1000.0
            march_penalty_arrear = 20.0
            april_principal_paid = 1000.0
            april_penalty_due = 40.0
            april_penalty_paid = 0.0
        elif member_id in ("NSF001", "NSF011", "NSF028"):
            # Paid principal + April late penalty in full
            april_penalty_due = source_penalty if source_penalty > 0 else (40.0 if member_id == "NSF011" else 0.0)
            april_penalty_paid = april_penalty_due
            april_principal_paid = principal_due
        elif paid_in_january:
            april_principal_paid = principal_due
            april_penalty_due = 0.0
        else:
            april_principal_paid = min(principal_due, received)
            april_penalty_due = source_penalty if member_id in ("NSF002", "NSF018", "NSF019", "NSF020", "NSF032", "NSF033") else 0.0

        ref = f"HIST-2024-04-R{row:02d}"
        status = "PAID" if (april_principal_paid >= principal_due and (april_penalty_due - april_penalty_paid) == 0) else ("PARTIAL" if (april_principal_paid > 0 or april_penalty_paid > 0) else "DUE")

        ledgers.append({
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "month": MONTH,
            "share_count": shares,
            "principal_due": principal_due,
            "principal_paid": april_principal_paid,
            "penalty_due": april_penalty_due,
            "penalty_paid": april_penalty_paid,
            "advance_applied": april_principal_paid if paid_in_january else 0.0,
            "excess_advance": 0.0,
            "status": status,
            "source_row": row,
            "comment": comment
        })

        if received <= 0:
            continue

        payment_ref = f"{ref}-P01"
        payments.append({
            "payment_ref": payment_ref,
            "source_ref": ref,
            "member_id": member_id,
            "member_name": name,
            "payment_date": value(ws.cell(row, 7).value),
            "total_amount": received,
            "gateway": gateway(ws.cell(row, 5).value),
            "receiver_source": value(ws.cell(row, 6).value),
            "comment": comment
        })

        if march_principal_arrear > 0:
            allocations.append({
                "payment_ref": payment_ref,
                "member_id": member_id,
                "target_month": "2024-03",
                "allocation_type": "PRINCIPAL",
                "amount": march_principal_arrear,
                "source_ref": ref
            })
        if march_penalty_arrear > 0:
            allocations.append({
                "payment_ref": payment_ref,
                "member_id": member_id,
                "target_month": "2024-03",
                "allocation_type": "PENALTY",
                "amount": march_penalty_arrear,
                "source_ref": ref
            })
        if april_principal_paid > 0 and not paid_in_january:
            allocations.append({
                "payment_ref": payment_ref,
                "member_id": member_id,
                "target_month": MONTH,
                "allocation_type": "PRINCIPAL",
                "amount": april_principal_paid,
                "source_ref": ref
            })
        if april_penalty_paid > 0:
            allocations.append({
                "payment_ref": payment_ref,
                "member_id": member_id,
                "target_month": MONTH,
                "allocation_type": "PENALTY",
                "amount": april_penalty_paid,
                "source_ref": ref
            })

    manifest = {"ledgers": ledgers, "payments": payments, "allocations": allocations}
    (OUT / "staging-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    summary = {
        "members": len(ledgers),
        "payments": len(payments),
        "allocations": len(allocations),
        "cash_received": sum(row["total_amount"] for row in payments),
        "april_principal_due": sum(row["principal_due"] - row["principal_paid"] for row in ledgers),
        "april_penalty_due": sum(row["penalty_due"] - row["penalty_paid"] for row in ledgers),
        "march_arrear_principal_collected": sum(row["amount"] for row in allocations if row["target_month"] == "2024-03" and row["allocation_type"] == "PRINCIPAL"),
        "march_arrear_penalty_collected": sum(row["amount"] for row in allocations if row["target_month"] == "2024-03" and row["allocation_type"] == "PENALTY"),
        "april_penalty_collected": sum(row["amount"] for row in allocations if row["target_month"] == "2024-04" and row["allocation_type"] == "PENALTY"),
    }
    (OUT / "verification-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
