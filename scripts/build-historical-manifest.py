"""Build private, review-only normalized manifests from historical workbook evidence.

This script never writes to MongoDB and does not infer financial allocation from
ambiguous comments or colour coding. It preserves source location and raw values
for administrator reconciliation before any posting implementation is approved.
"""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path("docs/import-historical-evidence")
OUT = ROOT / "normalized-manifests"
PAYMENT_FILES = sorted(path for path in ROOT.glob("*.xlsx") if path.name != "Investment.xlsx")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def number(value: Any) -> str:
    if value in (None, "", "-"):
        return ""
    if isinstance(value, (int, float)):
        return f"{float(value):.2f}"
    return text(value)


def period_for(path: Path, ws_title: str) -> str:
    match = re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December).*?(2024|2025)", f"{path.name} {ws_title}", re.I)
    if not match:
        match = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec).*?(2024|2025)", f"{path.name} {ws_title}", re.I)
    if not match:
        return "UNRESOLVED"
    month_text = match.group(1).lower()[:3]
    month = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}[month_text]
    return f"{match.group(2)}-{month:02d}"


def fill(cell: Any) -> str:
    color = cell.fill.fgColor
    return color.rgb or color.indexed or color.theme or ""


def normalized_gateway(value: Any) -> str:
    raw = text(value).lower()
    if raw in {"bkash", "b kash"}:
        return "bKash"
    if raw == "nagad":
        return "Nagad"
    if raw == "self":
        return "Cash"
    if raw in {"cash", "bank"}:
        return raw.title()
    if raw in {"", "select", "-"}:
        return ""
    return text(value)


def header_row(ws: Any) -> int:
    for row in range(1, min(ws.max_row or 80, 80) + 1):
        values = [text(ws.cell(row, col).value).lower() for col in range(1, min(ws.max_column or 20, 20) + 1)]
        if "member id" in values and any("name" in value for value in values):
            return row
    raise ValueError(f"No member header found in {ws.title}")


def header_map(ws: Any, row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, (ws.max_column or 20) + 1):
        label = text(ws.cell(row, col).value).strip().lower()
        if label and label not in headers:
            headers[label] = col
    return headers


def value(ws: Any, row: int, headers: dict[str, int], *names: str) -> Any:
    for name in names:
        if name.lower() in headers:
            return ws.cell(row, headers[name.lower()]).value
    return None


def cell_for(ws: Any, row: int, headers: dict[str, int], *names: str) -> Any | None:
    for name in names:
        if name.lower() in headers:
            return ws.cell(row, headers[name.lower()])
    return None


def load_sheet(path: Path, title: str) -> tuple[Any, Any]:
    formula = openpyxl.load_workbook(path, data_only=False)[title]
    values = openpyxl.load_workbook(path, data_only=True)[title]
    return formula, values


def extract_payments() -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    records: list[dict[str, str]] = []
    schemas: list[dict[str, str]] = []
    for path in PAYMENT_FILES:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
        for title in workbook.sheetnames:
            probe = workbook[title]
            row = header_row(probe)
            headers = header_map(probe, row)
            schemas.append({"file": path.name, "sheet": title, "period": period_for(path, title), "headers": ", ".join(headers.keys())})
            formula_ws, value_ws = load_sheet(path, title)
            formula_headers, value_headers = header_map(formula_ws, row), header_map(value_ws, row)
            for source_row in range(row + 1, (formula_ws.max_row or 0) + 1):
                member_id = text(value(value_ws, source_row, value_headers, "Member ID"))
                if not re.fullmatch(r"NSF\d+", member_id, re.I):
                    continue
                amount = value(value_ws, source_row, value_headers, "Paid Amount", "Amount")
                payment_date = value(value_ws, source_row, value_headers, "Payment Date", "Date of Payment")
                penalty_cell = cell_for(formula_ws, source_row, formula_headers, "Penalty", "Paid Penalty", "Due Penalty")
                due_cell = cell_for(formula_ws, source_row, formula_headers, "Due")
                total_received = number(amount)
                status = "PAYMENT_CANDIDATE" if total_received and total_received not in {"0.00", "0"} else "DUE_OR_UNPAID_REVIEW"
                records.append({
                    "source_file": path.name,
                    "source_sheet": title,
                    "source_row": str(source_row),
                    "period": period_for(path, title),
                    "member_id": member_id.upper(),
                    "member_name": text(value(value_ws, source_row, value_headers, "Member Name", "Name")),
                    "shares": number(value(value_ws, source_row, value_headers, "No. of Shares", "No. of Share")),
                    "scheduled_amount_raw": number(value(value_ws, source_row, value_headers, "Amount")),
                    "total_received_raw": total_received,
                    "payment_date": text(payment_date),
                    "gateway_raw": text(value(value_ws, source_row, value_headers, "Gateway", "Getway", "Gatetway")),
                    "gateway_normalized": normalized_gateway(value(value_ws, source_row, value_headers, "Gateway", "Getway", "Gatetway")),
                    "receiver_raw": text(value(value_ws, source_row, value_headers, "Receiver")),
                    "previous_due_raw": number(value(value_ws, source_row, value_headers, "Previous Due")),
                    "due_raw": number(value(value_ws, source_row, value_headers, "Due")),
                    "penalty_raw": number(value(value_ws, source_row, value_headers, "Penalty", "Due Penalty")),
                    "paid_penalty_raw": number(value(value_ws, source_row, value_headers, "Paid Penalty")),
                    "cashout_charge_raw": number(value(value_ws, source_row, value_headers, "CO Charge")),
                    "balance_raw": number(value(value_ws, source_row, value_headers, "Balance")),
                    "comment": text(value(value_ws, source_row, value_headers, "Comment")),
                    "penalty_fill": fill(penalty_cell) if penalty_cell else "",
                    "due_fill": fill(due_cell) if due_cell else "",
                    "manifest_status": status,
                    "posting_decision": "PENDING_RECONCILIATION",
                })
    seen: dict[tuple[str, ...], dict[str, str]] = {}
    for record in records:
        signature = tuple(record[key] for key in ("period", "member_id", "total_received_raw", "payment_date", "gateway_normalized", "receiver_raw", "due_raw", "penalty_raw", "paid_penalty_raw", "cashout_charge_raw", "balance_raw", "comment"))
        original = seen.get(signature)
        if original:
            record["manifest_status"] = "DUPLICATE_SOURCE_REVIEW"
            record["posting_decision"] = "EXCLUDE_UNTIL_DUPLICATE_RESOLVED"
        else:
            seen[signature] = record
    return records, schemas


def extract_investments() -> list[dict[str, str]]:
    path = ROOT / "Investment.xlsx"
    ws = openpyxl.load_workbook(path, data_only=True)["Investment"]
    headers = header_map(ws, 2)
    records: list[dict[str, str]] = []
    for row in range(3, (ws.max_row or 0) + 1):
        serial = text(value(ws, row, headers, "SL No."))
        if not serial:
            continue
        records.append({
            "source_file": path.name, "source_sheet": ws.title, "source_row": str(row), "source_serial": serial,
            "invoice_to": text(value(ws, row, headers, "Invoice to")), "invoice_no": text(value(ws, row, headers, "Invoice No")),
            "pay_date": text(value(ws, row, headers, "Pay Date")), "amount": number(value(ws, row, headers, "Amount")),
            "start_date": text(value(ws, row, headers, "Start Date")), "end_date": text(value(ws, row, headers, "End Date")),
            "project_name": text(value(ws, row, headers, "Project Name")), "duration": text(value(ws, row, headers, "Duration")),
            "status": text(value(ws, row, headers, "Project Status")), "organization": text(value(ws, row, headers, "Org")),
            "roi_raw": number(value(ws, row, headers, "ROI")), "gateway_raw": text(value(ws, row, headers, "Gateway")),
            "gateway_normalized": normalized_gateway(value(ws, row, headers, "Gateway")), "paid_by_wallet": number(value(ws, row, headers, "Paid By-Wallet (Re-invest)")),
            "paid_by_moin": number(value(ws, row, headers, "Paid By-Moin")), "paid_by_samrat": number(value(ws, row, headers, "Paid By-Samrat")),
            "sent_to": text(value(ws, row, headers, "Sent to")), "posting_decision": "PENDING_PROJECT_AND_CUSTODY_RECONCILIATION",
        })
    return records


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    columns = list(rows[0]) if rows else []
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader(); writer.writerows(rows)


def as_amount(value: str) -> float:
    try: return float(value)
    except (TypeError, ValueError): return 0.0


def write_report(payments: list[dict[str, str]], investments: list[dict[str, str]], schemas: list[dict[str, str]]) -> None:
    by_period: dict[str, list[dict[str, str]]] = defaultdict(list)
    for item in payments: by_period[item["period"]].append(item)
    lines = ["# Historical import manifest review", "", "This review-only manifest is generated from the private workbook evidence. It does not post, alter, or approve database records.", "", "## Payment workbook inventory", "", "| Period | Rows | Payment candidates | Received total | Due/unpaid review rows |", "|---|---:|---:|---:|---:|"]
    for period, rows in sorted(by_period.items()):
        candidates = [row for row in rows if row["manifest_status"] == "PAYMENT_CANDIDATE"]
        duplicates = [row for row in rows if row["manifest_status"] == "DUPLICATE_SOURCE_REVIEW"]
        lines.append(f"| {period} | {len(rows)} | {len(candidates)} | ৳{sum(as_amount(row['total_received_raw']) for row in candidates):,.2f} | {len(rows) - len(candidates) - len(duplicates)} |")
    duplicate_count = sum(1 for row in payments if row["manifest_status"] == "DUPLICATE_SOURCE_REVIEW")
    lines += ["", f"Detected {duplicate_count} potential duplicate source rows. They are retained for evidence but excluded from posting until resolved.", "", "## Detected schemas", "", "| File | Period | Columns |", "|---|---|---|"]
    for schema in schemas: lines.append(f"| {schema['file']} | {schema['period']} | {schema['headers']} |")
    gateway_counts = Counter(row["gateway_normalized"] or "Unspecified" for row in payments)
    lines += ["", "## Normalization rules applied", "", "- `Bkash`, `bkash`, and `BKash` are normalized to `bKash`.", "- `Nagad` and `nagad` are normalized to `Nagad`.", "- `Self` is normalized to `Cash`, per the supplied accounting interpretation.", "- Payment amount, dues, penalty, cash-out charge, balance, comments, original gateway, source row, and colour metadata are retained. No allocation is inferred from colour alone.", "", "## Required reconciliation before posting", "", "- Confirm one source member ID maps to one production member.", "- Resolve date/period conflicts, including advance payments recorded in a prior month.", "- Review all `DUE_OR_UNPAID_REVIEW` rows and color-coded penalty/due values with the accountant.", "- Reconcile each accepted payment to a valid receiver and a historical custody opening/transfer position.", "- Reconcile each investment to project, funding source, payment, return, and custody evidence.", "", "## Investment inventory", "", f"The investment manifest has {len(investments)} source rows. Every row remains pending project and custody reconciliation.", "", "## Files generated", "", "- `payment-manifest-review.csv`: one normalized source row per historical member/month row.", "- `investment-manifest-review.csv`: one normalized investment source row.", "- `manifest-analysis.md`: this reconciliation summary."]
    (OUT / "manifest-analysis.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payments, schemas = extract_payments()
    investments = extract_investments()
    write_csv(OUT / "payment-manifest-review.csv", payments)
    write_csv(OUT / "investment-manifest-review.csv", investments)
    write_report(payments, investments, schemas)
    print(f"Created {len(payments)} payment manifest rows and {len(investments)} investment manifest rows in {OUT}")


if __name__ == "__main__":
    main()
